import shutil

import openpyxl
import pytest

from conftest import ROOT
from rmbs import workbook as W


def make(tmp_path):
    for d in ("config", "out"):
        shutil.copytree(ROOT / d, tmp_path / d)
    for f in ("filing_locations.csv", W.WB):
        shutil.copy(ROOT / f, tmp_path / f)
    return tmp_path


def test_refresh_keeps_formulas_and_adds_shortfall_option(tmp_path):
    root = make(tmp_path)
    n0 = W.count_formulas(root / W.WB)
    assert W.main(["--config", str(root / "config" / "run.yaml"), "--no-recalc"]) == 0
    wb = openpyxl.load_workbook(root / W.WB)
    assert W.count_formulas(root / W.WB) >= n0
    h = wb["History"]
    assert [c.value for c in h[3]][:49] == W.SCHEMA
    assert h["AX3"].value == "chk: beg−prin−end" and h["AX4"].value.startswith("=IF(N4")
    isins = [wb["Deals"].cell(r, 1).value for r in range(2, 19)]
    assert len(set(isins)) == 17
    assert wb["Deals"]["AD1"].value == "IPD rule" and wb["Inputs"]["G20"].value == 1
    assert "BE11" in wb["Engine"]["AD11"].value and "Inputs!$G$20" in wb["Engine"]["AD11"].value
    assert "History!BC4" in wb["Inputs"]["G39"].value
    hipo = next(r for r in range(2, 19) if wb["Deals"].cell(r, 1).value == "ES0345672010")
    assert wb["Deals"].cell(hipo, 36).value.startswith("YES")
    assert W.scan_errors(root / W.WB) == []


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_recalc_zero_errors(tmp_path):
    root = make(tmp_path)
    assert W.main(["--config", str(root / "config" / "run.yaml")]) == 0
    wb = openpyxl.load_workbook(root / W.WB, data_only=True)
    assert isinstance(wb["Inputs"]["K9"].value, float) and wb["Inputs"]["K20"].value == pytest.approx(0, abs=1e-6)
